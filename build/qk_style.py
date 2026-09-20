"""Shared styling helpers for QuotaKit templates.

Design system: navy header band, teal accents, soft grey grid, blue input cells
with light-yellow fill (the classic 'edit these cells' convention), black formulas.
All fonts Arial. Formulas restricted to Excel-2007-era functions so the files
open cleanly in Google Sheets, Numbers and LibreOffice.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule, DataBarRule
from openpyxl.comments import Comment

NAVY = "1F2A44"
TEAL = "0E9F8E"
TEAL_LIGHT = "D9F2EF"
GREY_LINE = "D0D5DD"
GREY_FILL = "F4F6F8"
INPUT_FILL = "FFF8DC"      # light yellow: user edits these
INPUT_FONT = "0000FF"      # blue text for inputs
WHITE = "FFFFFF"
RED_FILL = "FDE2E1"
AMBER_FILL = "FFF1CC"
GREEN_FILL = "DDF3E4"

thin = Side(style="thin", color=GREY_LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

F_BASE = Font(name="Arial", size=10, color="1F2A44")
F_BOLD = Font(name="Arial", size=10, bold=True, color="1F2A44")
F_INPUT = Font(name="Arial", size=10, color=INPUT_FONT)
F_INPUT_BOLD = Font(name="Arial", size=10, bold=True, color=INPUT_FONT)
F_HEADER = Font(name="Arial", size=10, bold=True, color=WHITE)
F_TITLE = Font(name="Arial", size=18, bold=True, color=WHITE)
F_SUB = Font(name="Arial", size=10, italic=True, color="C9D3E0")
F_SECTION = Font(name="Arial", size=12, bold=True, color=TEAL)
F_NOTE = Font(name="Arial", size=9, italic=True, color="667085")
F_KPI = Font(name="Arial", size=20, bold=True, color=NAVY)
F_KPI_LABEL = Font(name="Arial", size=9, bold=True, color="667085")

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_TEAL = PatternFill("solid", fgColor=TEAL)
FILL_TEAL_LIGHT = PatternFill("solid", fgColor=TEAL_LIGHT)
FILL_GREY = PatternFill("solid", fgColor=GREY_FILL)
FILL_INPUT = PatternFill("solid", fgColor=INPUT_FILL)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

FMT_MONEY = '#,##0;(#,##0);"-"'
FMT_MONEY2 = '#,##0.00;(#,##0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_PCT0 = '0%;(0%);"-"'
FMT_INT = '#,##0;(#,##0);"-"'
FMT_DATE = 'dd-mmm-yyyy'


def title_band(ws, title, subtitle, ncols, row=1):
    """Navy title band across ncols, two rows tall."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.font = F_TITLE
    c.fill = FILL_NAVY
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s = ws.cell(row=row + 1, column=1, value=subtitle)
    s.font = F_SUB
    s.fill = FILL_NAVY
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = FILL_NAVY
        ws.cell(row=row + 1, column=col).fill = FILL_NAVY
    ws.row_dimensions[row].height = 30
    ws.row_dimensions[row + 1].height = 18


def section(ws, row, col, text, ncols=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_SECTION
    c.alignment = LEFT
    if ncols > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + ncols - 1)
    ws.row_dimensions[row].height = 22
    return c


def header_row(ws, row, col, headers, fill=FILL_TEAL, widths=None, height=28):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = F_HEADER
        c.fill = fill
        c.alignment = CENTER
        c.border = BORDER
        if widths and i < len(widths) and widths[i]:
            ws.column_dimensions[get_column_letter(col + i)].width = widths[i]
    ws.row_dimensions[row].height = height


def input_cell(ws, ref, value=None, fmt=None, bold=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = F_INPUT_BOLD if bold else F_INPUT
    c.fill = FILL_INPUT
    c.border = BORDER
    c.alignment = Alignment(horizontal="right" if isinstance(value, (int, float)) else "left", vertical="center")
    if fmt:
        c.number_format = fmt
    return c


def formula_cell(ws, ref, formula, fmt=None, bold=False, fill=None, align=None):
    c = ws[ref]
    c.value = formula
    c.font = F_BOLD if bold else F_BASE
    c.border = BORDER
    c.alignment = align or RIGHT
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def label_cell(ws, ref, text, bold=False, fill=None, align=None, font=None):
    c = ws[ref]
    c.value = text
    c.font = font or (F_BOLD if bold else F_BASE)
    c.alignment = align or LEFT
    c.border = BORDER
    if fill:
        c.fill = fill
    return c


def note(ws, ref, text, merge_to=None):
    c = ws[ref]
    c.value = text
    c.font = F_NOTE
    c.alignment = LEFT_TOP
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")
        import re
        r0 = int(re.sub(r"[A-Z]+", "", ref)); r1 = int(re.sub(r"[A-Z]+", "", merge_to))
        for rr in range(r0, r1 + 1):
            ws.row_dimensions[rr].height = 18
    return c


def kpi_tile(ws, top_left_col, row, label, formula, fmt, width_cols=2):
    """Two-row KPI tile: label on top (small), value below (big)."""
    from openpyxl.utils import column_index_from_string
    ci = column_index_from_string(top_left_col)
    ws.merge_cells(start_row=row, start_column=ci, end_row=row, end_column=ci + width_cols - 1)
    ws.merge_cells(start_row=row + 1, start_column=ci, end_row=row + 1, end_column=ci + width_cols - 1)
    l = ws.cell(row=row, column=ci, value=label)
    l.font = F_KPI_LABEL
    l.alignment = CENTER
    l.fill = FILL_TEAL_LIGHT
    v = ws.cell(row=row + 1, column=ci, value=formula)
    v.font = F_KPI
    v.alignment = CENTER
    v.number_format = fmt
    v.fill = FILL_TEAL_LIGHT
    for cc in range(ci, ci + width_cols):
        ws.cell(row=row, column=cc).fill = FILL_TEAL_LIGHT
        ws.cell(row=row + 1, column=cc).fill = FILL_TEAL_LIGHT
        ws.cell(row=row, column=cc).border = BORDER
        ws.cell(row=row + 1, column=cc).border = BORDER
    ws.row_dimensions[row].height = 16
    ws.row_dimensions[row + 1].height = 34


def legend(ws, row, col, ncols=6):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + ncols - 1)
    c = ws.cell(row=row, column=col,
                value="How to use: edit only the yellow cells (blue text). Everything else calculates automatically. "
                      "Do not insert rows inside tables; extend by filling the next empty row.")
    c.font = F_NOTE
    c.alignment = LEFT
    ws.row_dimensions[row].height = 26


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def freeze(ws, ref):
    ws.freeze_panes = ref


def hide_gridlines(ws):
    ws.sheet_view.showGridLines = False


def add_list_validation(ws, rng, formula, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "Pick from the list"
    dv.error = "Please choose a value from the dropdown."
    if prompt:
        dv.promptTitle = "Tip"
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def instructions_sheet(wb, product_name, steps, tabs, tips, first=True):
    """Standard 'Start Here' tab."""
    ws = wb.create_sheet("Start Here", 0 if first else None)
    hide_gridlines(ws)
    set_widths(ws, {"A": 3, "B": 26, "C": 90, "D": 3})
    title_band(ws, product_name, "QuotaKit  ·  Start Here  ·  Read this tab once, then go to Settings.", 4)
    r = 4
    section(ws, r, 2, "Quick start", 2)
    r += 1
    for i, s in enumerate(steps, 1):
        label_cell(ws, f"B{r}", f"Step {i}", bold=True, fill=FILL_GREY)
        label_cell(ws, f"C{r}", s)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    section(ws, r, 2, "What each tab does", 2)
    r += 1
    for name, desc in tabs:
        label_cell(ws, f"B{r}", name, bold=True, fill=FILL_GREY)
        label_cell(ws, f"C{r}", desc)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    section(ws, r, 2, "Colour code", 2)
    r += 1
    input_cell(ws, f"B{r}", "Yellow cell, blue text")
    label_cell(ws, f"C{r}", "Your inputs. Edit freely.")
    r += 1
    label_cell(ws, f"B{r}", "White cell, black text", bold=False)
    label_cell(ws, f"C{r}", "Formulas. Leave alone (they are unprotected so you can adapt them if you want).")
    r += 1
    label_cell(ws, f"B{r}", "Teal tiles", fill=FILL_TEAL_LIGHT)
    label_cell(ws, f"C{r}", "Headline numbers that update as you type.")
    r += 2
    section(ws, r, 2, "Tips", 2)
    r += 1
    for t in tips:
        label_cell(ws, f"B{r}", "•", align=CENTER)
        label_cell(ws, f"C{r}", t)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    note(ws, f"B{r}", "Works in Microsoft Excel 2010+ and Google Sheets (File > Import > Upload). Built with formulas only, no macros. "
                     "Personal and small-business licence: use for your own team or clients; please do not resell or redistribute the file. "
                     "Questions or a feature request? Message QuotaKit through the shop.", merge_to=f"C{r}")
    ws.row_dimensions[r].height = 48
    return ws
