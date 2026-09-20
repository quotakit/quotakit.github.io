"""QuotaKit — Territory & Account Coverage Planner (Excel + Google Sheets)."""
import sys, random
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule, DataBarRule, CellIsRule
from openpyxl.styles import PatternFill, Font
from qk_style import *

SAMPLE = "--sample" in sys.argv
OUT = sys.argv[-1] if sys.argv[-1].endswith(".xlsx") else "out.xlsx"
NACC = 300
NREP = 25
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

wb = Workbook(); wb.remove(wb.active)
instructions_sheet(
    wb, "Territory & Account Coverage Planner",
    steps=[
        "Settings tab: set the visit frequency for A, B and C accounts, the year, and how many months are complete.",
        "Reps tab: list your field reps (up to 25).",
        "Accounts tab: one row per account — name, area, rep, tier, potential. Then log actual visits per month in the yellow month columns.",
        "Coverage, gap and 'missed' flags calculate per account. Rep Summary rolls it up by rep.",
        "Dashboard shows overall coverage, A-account coverage, planned vs actual by month and the rep league table.",
    ],
    tabs=[
        ("Settings", "Visit frequency per tier, year, months completed, currency."),
        ("Reps", "Rep roster used by the Accounts dropdown."),
        ("Accounts", "Account list with monthly actual visits (up to 300 accounts)."),
        ("Rep Summary", "Per rep: accounts by tier, planned vs actual visits, coverage %, potential covered."),
        ("Dashboard", "Tiles, monthly planned vs actual chart, tier coverage, rep league table."),
    ],
    tips=[
        "Tier = how important the account is. A = highest potential, visited most often. Set frequencies that your team can actually deliver.",
        "'Months completed' controls the planned-visit denominator so coverage % is fair mid-year. Update it monthly.",
        "'Missed' = an A or B account with zero visits in the latest completed month. Start every review there.",
        "Potential value is optional, but with it the dashboard tells you how much potential your uncovered accounts represent.",
    ],
)

# ---------------------------------------------------------------- Settings
ws = wb.create_sheet("Settings"); hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 30, "C": 16, "D": 3, "E": 56})
title_band(ws, "Settings", "Edit the yellow cells.", 5)
section(ws, 4, 2, "Visit frequency (visits per account per month)")
header_row(ws, 5, 2, ["Tier", "Visits / month"])
for i, (t, v) in enumerate([("A", 4), ("B", 2), ("C", 1)]):
    r = 6 + i
    label_cell(ws, f"B{r}", t, bold=True, fill=FILL_GREY, align=CENTER); input_cell(ws, f"C{r}", v, fmt="0")
section(ws, 10, 2, "General")
label_cell(ws, "B11", "Year", bold=True, fill=FILL_GREY); input_cell(ws, "C11", 2026, fmt="0")
label_cell(ws, "B12", "Months completed (1–12)", bold=True, fill=FILL_GREY); input_cell(ws, "C12", 8 if SAMPLE else 1, fmt="0")
add_list_validation(ws, "C12", '"1,2,3,4,5,6,7,8,9,10,11,12"')
label_cell(ws, "B13", "Currency symbol / code", bold=True, fill=FILL_GREY); input_cell(ws, "C13", "$")
note(ws, "E5", "Planned visits = tier frequency × months completed. Coverage % = actual ÷ planned.", merge_to="E8")
note(ws, "E12", "Set this to the last month you have fully logged. It drives planned visits and the 'missed' flag.")

# ---------------------------------------------------------------- Reps
ws = wb.create_sheet("Reps"); hide_gridlines(ws)
set_widths(ws, {"A": 24, "B": 20, "C": 3, "D": 50})
title_band(ws, "Reps", "Up to 25 field reps.", 4)
header_row(ws, 4, 1, ["Rep", "Region"])
reps = [("Aarav K.", "West"), ("Bhavna S.", "West"), ("Chirag M.", "North"), ("Divya R.", "North"), ("Eshan P.", "South"), ("Farida D.", "South"), ("Gaurav V.", "East"), ("Hina G.", "East")]
for i in range(NREP):
    r = 5 + i
    if SAMPLE and i < len(reps):
        input_cell(ws, f"A{r}", reps[i][0]); input_cell(ws, f"B{r}", reps[i][1])
    else:
        input_cell(ws, f"A{r}"); input_cell(ws, f"B{r}")

# ---------------------------------------------------------------- Accounts
ws = wb.create_sheet("Accounts"); hide_gridlines(ws)
heads = ["Account", "Area / city", "Rep", "Tier", "Potential (annual)"] + MONTHS + ["Planned YTD", "Actual YTD", "Coverage %", "Gap", "Missed?"]
widths = [26, 14, 14, 7, 14] + [6] * 12 + [11, 11, 11, 8, 10]
title_band(ws, "Accounts", "One row per account. Yellow = you type (details + monthly actual visits). Grey = calculates.", len(heads))
header_row(ws, 4, 1, heads, widths=widths)
for c in range(18, 23):
    ws.cell(row=4, column=c).fill = FILL_NAVY
add_list_validation(ws, f"C5:C{4+NACC}", f"=Reps!$A$5:$A${4+NREP}")
add_list_validation(ws, f"D5:D{4+NACC}", '"A,B,C"')
random.seed(21)
names = ["Apex Clinic", "Bluebird Pharmacy", "Cedar Hospital", "Delta Diagnostics", "Evergreen Care", "Fortune Medical", "Galaxy Labs", "Harbor Health",
         "Iris Wellness", "Juniper Clinic", "Keystone Hospital", "Lakeside Care", "Maple Pharmacy", "Nova Medical", "Orchid Clinic", "Pinnacle Health",
         "Quartz Labs", "Riverside Hospital", "Summit Care", "Trident Pharmacy", "Unity Clinic", "Vista Medical", "Willow Health", "Zenith Hospital"]
areas = ["Andheri", "Bandra", "Dadar", "Thane", "Navi Mumbai", "Powai", "Borivali", "Colaba"]
for i in range(NACC):
    r = 5 + i
    if SAMPLE and i < 80:
        tier = random.choices("ABC", weights=[25, 40, 35])[0]
        freqs = {"A": 4, "B": 2, "C": 1}
        input_cell(ws, f"A{r}", f"{random.choice(names)} {i+1}"); input_cell(ws, f"B{r}", random.choice(areas))
        input_cell(ws, f"C{r}", random.choice(reps)[0]); input_cell(ws, f"D{r}", tier)
        input_cell(ws, f"E{r}", random.choice([20000, 45000, 80000, 120000, 250000]) if tier != "C" else random.choice([5000, 12000, 20000]), fmt=FMT_MONEY)
        for m in range(12):
            v = max(0, freqs[tier] + random.choice([-2, -1, -1, 0, 0, 0, 1])) if m < 8 else None
            if m < 8 and random.random() < 0.08: v = 0
            input_cell(ws, f"{L(6+m)}{r}", v, fmt="0")
    else:
        input_cell(ws, f"A{r}"); input_cell(ws, f"B{r}"); input_cell(ws, f"C{r}"); input_cell(ws, f"D{r}"); input_cell(ws, f"E{r}", fmt=FMT_MONEY)
        for m in range(12):
            input_cell(ws, f"{L(6+m)}{r}", fmt="0")
    formula_cell(ws, f"R{r}", f'=IF(D{r}="","",IFERROR(INDEX(Settings!$C$6:$C$8,MATCH(D{r},Settings!$B$6:$B$8,0)),0)*Settings!$C$12)', fmt="0", fill=FILL_GREY)
    formula_cell(ws, f"S{r}", f'=IF(D{r}="","",SUM(F{r}:Q{r}))', fmt="0", fill=FILL_GREY)
    formula_cell(ws, f"T{r}", f'=IF(D{r}="","",IF(R{r}>0,S{r}/R{r},0))', fmt=FMT_PCT0, fill=FILL_GREY)
    formula_cell(ws, f"U{r}", f'=IF(D{r}="","",MAX(0,R{r}-S{r}))', fmt="0", fill=FILL_GREY)
    formula_cell(ws, f"V{r}", f'=IF(D{r}="","",IF(AND(OR(D{r}="A",D{r}="B"),N(INDEX(F{r}:Q{r},1,Settings!$C$12))=0),"MISSED",""))', align=CENTER, fill=FILL_GREY)
ws.conditional_formatting.add(f"V5:V{4+NACC}", CellIsRule(operator="equal", formula=['"MISSED"'], fill=PatternFill("solid", fgColor=RED_FILL), font=Font(name="Arial", size=10, bold=True, color="B42318")))
ws.conditional_formatting.add(f"T5:T{4+NACC}", FormulaRule(formula=['AND($T5<>"",$T5<0.7)'], fill=PatternFill("solid", fgColor=RED_FILL)))
ws.conditional_formatting.add(f"T5:T{4+NACC}", FormulaRule(formula=['AND($T5<>"",$T5>=0.9)'], fill=PatternFill("solid", fgColor=GREEN_FILL)))
freeze(ws, "F5")
ws.auto_filter.ref = f"A4:V{4+NACC}"

# ---------------------------------------------------------------- Rep Summary
ws = wb.create_sheet("Rep Summary"); hide_gridlines(ws)
set_widths(ws, {"A": 24, "B": 12, "C": 8, "D": 8, "E": 8, "F": 12, "G": 12, "H": 12, "I": 10, "J": 16, "K": 16})
title_band(ws, "Rep Summary", "Coverage rolled up per rep. No inputs here.", 11)
header_row(ws, 4, 1, ["Rep", "Accounts", "A", "B", "C", "Planned YTD", "Actual YTD", "Coverage %", "Missed", "Potential (total)", "Potential in missed"])
A = "Accounts"
def rng(c): return f"{A}!${c}$5:${c}${4+NACC}"
for i in range(NREP):
    r = 5 + i
    formula_cell(ws, f"A{r}", f'=IF(Reps!A{r}="","",Reps!A{r})', align=LEFT, fill=FILL_GREY)
    formula_cell(ws, f"B{r}", f'=IF(A{r}="","",COUNTIF({rng("C")},A{r}))', fmt="0")
    for j, t in enumerate("ABC"):
        formula_cell(ws, f"{L(3+j)}{r}", f'=IF($A{r}="","",COUNTIFS({rng("C")},$A{r},{rng("D")},"{t}"))', fmt="0")
    formula_cell(ws, f"F{r}", f'=IF(A{r}="","",SUMIF({rng("C")},A{r},{rng("R")}))', fmt="0")
    formula_cell(ws, f"G{r}", f'=IF(A{r}="","",SUMIF({rng("C")},A{r},{rng("S")}))', fmt="0")
    formula_cell(ws, f"H{r}", f'=IF(A{r}="","",IF(F{r}>0,G{r}/F{r},0))', fmt=FMT_PCT0, bold=True)
    formula_cell(ws, f"I{r}", f'=IF(A{r}="","",COUNTIFS({rng("C")},A{r},{rng("V")},"MISSED"))', fmt="0")
    formula_cell(ws, f"J{r}", f'=IF(A{r}="","",SUMIF({rng("C")},A{r},{rng("E")}))', fmt=FMT_MONEY)
    formula_cell(ws, f"K{r}", f'=IF(A{r}="","",SUMIFS({rng("E")},{rng("C")},A{r},{rng("V")},"MISSED"))', fmt=FMT_MONEY)
tr = 5 + NREP
label_cell(ws, f"A{tr}", "Team", bold=True, fill=FILL_TEAL_LIGHT)
for c in "BCDEFGIJK":
    formula_cell(ws, f"{c}{tr}", f"=SUM({c}5:{c}{tr-1})", fmt=FMT_MONEY if c in "JK" else "0", bold=True, fill=FILL_TEAL_LIGHT)
formula_cell(ws, f"H{tr}", f"=IF(F{tr}>0,G{tr}/F{tr},0)", fmt=FMT_PCT0, bold=True, fill=FILL_TEAL_LIGHT)
ws.conditional_formatting.add(f"H5:H{4+NREP}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="0E9F8E"))
freeze(ws, "B5")

# ---------------------------------------------------------------- Dashboard
ws = wb.create_sheet("Dashboard"); hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 3, "I": 24, "J": 12, "K": 12, "L": 12, "M": 3})
title_band(ws, "Coverage Dashboard", "Year-to-date coverage for the months completed. No inputs here.", 13)
RS = "'Rep Summary'"
kpi_tile(ws, "B", 4, "ACCOUNTS", f'=COUNTA({rng("A")})', "0", 1)
kpi_tile(ws, "C", 4, "COVERAGE YTD", f"={RS}!H{tr}", FMT_PCT0, 1)
kpi_tile(ws, "D", 4, "A-ACCOUNT COVERAGE", f'=IF(SUMIF({rng("D")},"A",{rng("R")})>0,SUMIF({rng("D")},"A",{rng("S")})/SUMIF({rng("D")},"A",{rng("R")}),0)', FMT_PCT0, 2)
kpi_tile(ws, "F", 4, "MISSED (A/B, LAST MONTH)", f'=COUNTIF({rng("V")},"MISSED")', "0", 2)
kpi_tile(ws, "I", 4, "POTENTIAL IN MISSED ACCOUNTS", f'=SUMIF({rng("V")},"MISSED",{rng("E")})', FMT_MONEY, 2)
kpi_tile(ws, "K", 4, "VISIT GAP (VISITS)", f'=SUM({rng("U")})', "0", 2)
section(ws, 7, 2, "Planned vs actual visits by month", 6)
header_row(ws, 8, 2, ["Month", "Planned", "Actual", "Coverage %"])
PLAN_PM = f'(COUNTIF({rng("D")},"A")*Settings!$C$6+COUNTIF({rng("D")},"B")*Settings!$C$7+COUNTIF({rng("D")},"C")*Settings!$C$8)'
for m in range(12):
    r = 9 + m
    label_cell(ws, f"B{r}", MONTHS[m], bold=True, fill=FILL_GREY)
    formula_cell(ws, f"C{r}", f'=IF({m+1}<=Settings!$C$12,{PLAN_PM},0)', fmt="0")
    formula_cell(ws, f"D{r}", f'=IF({m+1}<=Settings!$C$12,SUM({rng(L(6+m))}),0)', fmt="0")
    formula_cell(ws, f"E{r}", f'=IF(C{r}>0,D{r}/C{r},"")', fmt=FMT_PCT0)
ch = BarChart(); ch.type = "col"; ch.title = "Planned vs actual visits"; ch.height, ch.width = 8, 15
ch.add_data(Reference(ws, min_col=3, min_row=8, max_col=4, max_row=20), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=9, max_row=20))
ch.series[0].graphicalProperties.solidFill = "C9D3E0"; ch.series[1].graphicalProperties.solidFill = "0E9F8E"
ws.add_chart(ch, "B22")
section(ws, 7, 9, "Coverage by tier", 4)
header_row(ws, 8, 9, ["Tier", "Accounts", "Planned", "Coverage %"])
for i, t in enumerate("ABC"):
    r = 9 + i
    label_cell(ws, f"I{r}", f"Tier {t}", bold=True, fill=FILL_GREY)
    formula_cell(ws, f"J{r}", f'=COUNTIF({rng("D")},"{t}")', fmt="0")
    formula_cell(ws, f"K{r}", f'=SUMIF({rng("D")},"{t}",{rng("R")})', fmt="0")
    formula_cell(ws, f"L{r}", f'=IF(K{r}>0,SUMIF({rng("D")},"{t}",{rng("S")})/K{r},0)', fmt=FMT_PCT0, bold=True)
section(ws, 13, 9, "Rep league table", 4)
header_row(ws, 14, 9, ["Rep", "Coverage %", "Missed", "Accounts"])
# rank by coverage using helper in Rep Summary col M
rs = wb["Rep Summary"]
rs["M4"] = "rank key"; rs["M4"].font = F_NOTE
for i in range(NREP):
    r = 5 + i
    formula_cell(rs, f"M{r}", f'=IF(A{r}="",-1,N(H{r})+ROW()/1000000)', fmt="0.000000"); rs[f"M{r}"].font = F_NOTE
KEY = f"{RS}!$M$5:$M${4+NREP}"
for k in range(10):
    r = 15 + k
    formula_cell(ws, f"I{r}", f'=IFERROR(IF(LARGE({KEY},{k+1})<0,"",INDEX({RS}!$A$5:$A${4+NREP},MATCH(LARGE({KEY},{k+1}),{KEY},0))),"")', align=LEFT)
    formula_cell(ws, f"J{r}", f'=IF(I{r}="","",INDEX({RS}!$H$5:$H${4+NREP},MATCH(I{r},{RS}!$A$5:$A${4+NREP},0)))', fmt=FMT_PCT0, bold=True)
    formula_cell(ws, f"K{r}", f'=IF(I{r}="","",INDEX({RS}!$I$5:$I${4+NREP},MATCH(I{r},{RS}!$A$5:$A${4+NREP},0)))', fmt="0")
    formula_cell(ws, f"L{r}", f'=IF(I{r}="","",INDEX({RS}!$B$5:$B${4+NREP},MATCH(I{r},{RS}!$A$5:$A${4+NREP},0)))', fmt="0")
ws.conditional_formatting.add("J15:J24", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="0E9F8E"))
note(ws, "I26", "Coverage % = actual visits ÷ planned visits for the completed months. Missed = A/B accounts with zero visits in the latest completed month.", merge_to="L28")

for s in wb.worksheets:
    s.sheet_properties.tabColor = TEAL if s.title in ("Settings", "Reps", "Accounts") else NAVY
wb["Start Here"].sheet_properties.tabColor = "F2B134"
wb.save(OUT); print("saved", OUT)
