"""QuotaKit — Sales Incentive Plan Simulator (Excel + Google Sheets)."""
import sys, random
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter as L
from qk_style import *

SAMPLE = "--sample" in sys.argv
OUT = sys.argv[-1] if sys.argv[-1].endswith(".xlsx") else "out.xlsx"
NREPS = 100
NK = 8  # knots
SCEN = ["Base", "Upside", "Downside"]

wb = Workbook(); wb.remove(wb.active)

instructions_sheet(
    wb, "Sales Incentive Plan Simulator",
    steps=[
        "Plan Design tab: set the payout curve (attainment → % of target incentive), curve type (step or linear), cap and optional kicker.",
        "Team tab: list reps with their target and target incentive. Enter Base actuals; Upside/Downside pre-fill from the levers (overwrite if you have real numbers).",
        "Simulation tab calculates attainment, payout % and payout amount for every rep under all three scenarios.",
        "Payout Curve tab shows the curve you designed and lets you test any attainment value.",
        "Dashboard compares plan cost, cost-to-sales ratio and attainment distribution across scenarios.",
    ],
    tabs=[
        ("Plan Design", "Currency, period, curve type, up to 8 curve points (attainment %, payout % of target incentive), cap, kicker, scenario levers."),
        ("Team", "Rep roster (up to 100): target, target incentive, actual sales under three scenarios."),
        ("Simulation", "The engine. Attainment, payout %, payout amount, kicker and total per rep per scenario."),
        ("Payout Curve", "Curve table and chart from 0% to 200% attainment, plus a what-if calculator."),
        ("Dashboard", "Scenario comparison, attainment distribution and cost summary."),
    ],
    tips=[
        "Target incentive = what a rep earns at exactly 100% attainment. Payout % is expressed against that amount.",
        "Step curve: the payout % jumps at each point (slabs). Linear curve: payout rises smoothly between points and flattens after the last one.",
        "The first point is your threshold — below it nobody is paid. The cap limits payout % no matter how high attainment goes.",
        "Use the levers to stress-test: what does the plan cost if everyone lands 10% higher than plan? Is the cost-to-sales ratio still acceptable?",
    ],
)

# ---------------------------------------------------------------- Plan Design
ws = wb.create_sheet("Plan Design"); hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 38, "C": 18, "D": 20, "E": 3, "F": 64})
title_band(ws, "Plan Design", "Edit the yellow cells. The Simulation and Dashboard read from here.", 6)
section(ws, 4, 2, "General")
label_cell(ws, "B5", "Currency symbol / code", bold=True, fill=FILL_GREY); input_cell(ws, "C5", "$")
label_cell(ws, "B6", "Period label", bold=True, fill=FILL_GREY); input_cell(ws, "C6", "Q3 2026")
label_cell(ws, "B7", "Curve type", bold=True, fill=FILL_GREY); input_cell(ws, "C7", "Linear")
add_list_validation(ws, "C7", '"Step,Linear"', "Step = slab jumps. Linear = smooth between points.")
note(ws, "F7", "Step: payout % equals the value at the highest point reached. Linear: payout % is interpolated between points.")

section(ws, 9, 2, "Payout curve (attainment → payout as % of target incentive)")
header_row(ws, 10, 2, ["Point", "Attainment ≥", "Payout % of TI"])
knots = [(0.80, 0.50), (0.90, 0.75), (1.00, 1.00), (1.10, 1.25), (1.20, 1.50), (1.30, 1.75), (None, None), (None, None)]
for i, (a, p) in enumerate(knots, 1):
    r = 10 + i
    label_cell(ws, f"B{r}", f"Point {i}", bold=True, fill=FILL_GREY)
    input_cell(ws, f"C{r}", a, fmt=FMT_PCT0); input_cell(ws, f"D{r}", p, fmt=FMT_PCT0)
note(ws, "F10", "Points must increase down the table. Point 1 is the threshold: below it payout is 0. "
                "Leave unused points blank. Payout % can exceed 100% (accelerators).", merge_to="F13")
label_cell(ws, "B19", "Active points", bold=True, fill=FILL_GREY); formula_cell(ws, "C19", "=COUNT(C11:C18)", fmt="0")
label_cell(ws, "B20", "Cap: maximum payout % of TI", bold=True, fill=FILL_GREY); input_cell(ws, "C20", 2.0, fmt=FMT_PCT0)
note(ws, "F20", "Payout % is never higher than the cap. Set a high number (e.g. 999%) for an uncapped plan.")

section(ws, 22, 2, "Kicker (optional flat amount)")
label_cell(ws, "B23", "Kicker amount (0 = none)", bold=True, fill=FILL_GREY); input_cell(ws, "C23", 1000, fmt=FMT_MONEY)
label_cell(ws, "B24", "Paid at attainment ≥", bold=True, fill=FILL_GREY); input_cell(ws, "C24", 1.2, fmt=FMT_PCT0)

section(ws, 26, 2, "Scenario levers (pre-fill Upside / Downside actuals from Base)")
label_cell(ws, "B27", "Upside: Base actual ×", bold=True, fill=FILL_GREY); input_cell(ws, "C27", 1.10, fmt="0.00")
label_cell(ws, "B28", "Downside: Base actual ×", bold=True, fill=FILL_GREY); input_cell(ws, "C28", 0.90, fmt="0.00")
note(ws, "F27", "Team tab pre-fills Upside and Downside actuals as Base × these factors. Overwrite any cell there with real numbers if you have them.", merge_to="F28")

A_RNG, P_RNG, NPT, CAP, CT = "'Plan Design'!$C$11:$C$18", "'Plan Design'!$D$11:$D$18", "'Plan Design'!$C$19", "'Plan Design'!$C$20", "'Plan Design'!$C$7"
KICK, KICK_AT = "'Plan Design'!$C$23", "'Plan Design'!$C$24"

def payout_formula(att):
    """Payout % of TI for an attainment cell ref (string)."""
    k = f"MATCH({att},INDEX({A_RNG},1):INDEX({A_RNG},{NPT}),1)"
    step = f"INDEX({P_RNG},{k})"
    lin = (f"IF({k}>={NPT},INDEX({P_RNG},{NPT}),"
           f"INDEX({P_RNG},{k})+({att}-INDEX({A_RNG},{k}))/(INDEX({A_RNG},{k}+1)-INDEX({A_RNG},{k}))*(INDEX({P_RNG},{k}+1)-INDEX({P_RNG},{k})))")
    return f"IFERROR(MIN({CAP},IF({CT}=\"Step\",{step},{lin})),0)"

# ---------------------------------------------------------------- Team
ws = wb.create_sheet("Team"); hide_gridlines(ws)
set_widths(ws, {"A": 24, "B": 18, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16, "H": 3, "I": 50})
title_band(ws, "Team", "Up to 100 reps. Base actuals are inputs; Upside/Downside pre-fill from the levers on Plan Design.", 9)
header_row(ws, 4, 1, ["Rep", "Territory / role", "Target", "Target incentive", "Actual — Base", "Actual — Upside", "Actual — Downside"])
random.seed(11)
first = ["Aarav", "Bhavna", "Chirag", "Divya", "Eshan", "Farida", "Gaurav", "Hina", "Ishaan", "Jaya", "Kabir", "Leela", "Manav", "Neha",
         "Om", "Priya", "Qasim", "Ritu", "Sahil", "Tara", "Uday", "Vani", "Wasim", "Yash", "Zara"]
terr = ["Mumbai", "Delhi", "Bengaluru", "Chennai", "Pune", "Kolkata", "Hyderabad", "Ahmedabad"]
for i in range(NREPS):
    r = 5 + i
    if SAMPLE and i < 25:
        tgt = random.choice([400000, 500000, 600000, 750000])
        input_cell(ws, f"A{r}", f"{first[i]} {random.choice('KMRSPDGV')}.")
        input_cell(ws, f"B{r}", random.choice(terr))
        input_cell(ws, f"C{r}", tgt, fmt=FMT_MONEY)
        input_cell(ws, f"D{r}", int(tgt * 0.05), fmt=FMT_MONEY)
        input_cell(ws, f"E{r}", int(tgt * random.uniform(0.65, 1.35) / 1000) * 1000, fmt=FMT_MONEY)
    else:
        input_cell(ws, f"A{r}"); input_cell(ws, f"B{r}")
        input_cell(ws, f"C{r}", fmt=FMT_MONEY); input_cell(ws, f"D{r}", fmt=FMT_MONEY); input_cell(ws, f"E{r}", fmt=FMT_MONEY)
    input_cell(ws, f"F{r}", f'=IF(E{r}="","",ROUND(E{r}*\'Plan Design\'!$C$27,0))', fmt=FMT_MONEY)
    input_cell(ws, f"G{r}", f'=IF(E{r}="","",ROUND(E{r}*\'Plan Design\'!$C$28,0))', fmt=FMT_MONEY)
note(ws, "I4", "Target incentive (TI) = the rep's payout at exactly 100% attainment. Upside/Downside columns hold a pre-fill formula — type over it to use real numbers.", merge_to="I7")
freeze(ws, "B5")

# ---------------------------------------------------------------- Simulation
ws = wb.create_sheet("Simulation"); hide_gridlines(ws)
set_widths(ws, {"A": 24, "B": 14, "C": 14})
cols = ["Rep", "Target", "Target incentive"]
for sc in SCEN:
    cols += [f"{sc}: actual", f"{sc}: attainment", f"{sc}: payout %", f"{sc}: payout", f"{sc}: kicker", f"{sc}: total"]
title_band(ws, "Simulation", "Payout engine — no inputs on this tab. One block of six columns per scenario.", len(cols))
header_row(ws, 4, 1, cols)
for c in range(4, len(cols) + 1):
    ws.column_dimensions[L(c)].width = 13
for si, sc in enumerate(SCEN):
    c0 = 4 + si * 6
    ws.cell(row=3, column=c0, value=sc + " scenario").font = F_SECTION
    for c in range(c0, c0 + 6):
        ws.cell(row=4, column=c).fill = [FILL_TEAL, FILL_NAVY, PatternFill("solid", fgColor="4B5D80")][si]
for i in range(NREPS):
    r = 5 + i
    formula_cell(ws, f"A{r}", f'=IF(Team!A{r}="","",Team!A{r})', align=LEFT)
    formula_cell(ws, f"B{r}", f'=IF($A{r}="","",Team!C{r})', fmt=FMT_MONEY)
    formula_cell(ws, f"C{r}", f'=IF($A{r}="","",Team!D{r})', fmt=FMT_MONEY)
    for si, sc in enumerate(SCEN):
        c0 = 4 + si * 6
        act, att, pp, pay, kk, tot = [L(c0 + j) for j in range(6)]
        src = ["E", "F", "G"][si]
        formula_cell(ws, f"{act}{r}", f'=IF($A{r}="","",N(Team!{src}{r}))', fmt=FMT_MONEY)
        formula_cell(ws, f"{att}{r}", f'=IF($A{r}="","",IF(N($B{r})>0,{act}{r}/$B{r},0))', fmt=FMT_PCT)
        formula_cell(ws, f"{pp}{r}", f'=IF($A{r}="","",{payout_formula(att + str(r))})', fmt=FMT_PCT)
        formula_cell(ws, f"{pay}{r}", f'=IF($A{r}="","",ROUND({pp}{r}*N($C{r}),0))', fmt=FMT_MONEY)
        formula_cell(ws, f"{kk}{r}", f'=IF($A{r}="","",IF(AND({KICK}>0,{att}{r}>={KICK_AT}),{KICK},0))', fmt=FMT_MONEY)
        formula_cell(ws, f"{tot}{r}", f'=IF($A{r}="","",{pay}{r}+{kk}{r})', fmt=FMT_MONEY, bold=True)
tr = 5 + NREPS
label_cell(ws, f"A{tr}", "Total", bold=True, fill=FILL_GREY)
for c in range(2, len(cols) + 1):
    col = L(c)
    if cols[c - 1].endswith("attainment") or cols[c - 1].endswith("payout %"):
        continue
    formula_cell(ws, f"{col}{tr}", f"=SUM({col}5:{col}{tr-1})", fmt=FMT_MONEY, bold=True, fill=FILL_GREY)
freeze(ws, "D5")

# ---------------------------------------------------------------- Payout Curve
ws = wb.create_sheet("Payout Curve"); hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 16, "C": 16, "D": 3, "E": 30, "F": 18, "G": 3})
title_band(ws, "Payout Curve", "Your curve from 0% to 200% attainment, and a what-if calculator.", 7)
header_row(ws, 4, 2, ["Attainment", "Payout % of TI"])
for i in range(41):
    r = 5 + i
    label_cell(ws, f"B{r}", i * 0.05, align=RIGHT); ws[f"B{r}"].number_format = FMT_PCT0
    formula_cell(ws, f"C{r}", f"={payout_formula('B' + str(r))}", fmt=FMT_PCT0)
ch = LineChart(); ch.title = "Payout curve"; ch.y_axis.title = "Payout % of target incentive"; ch.x_axis.title = "Attainment"
ch.height, ch.width = 9, 16
ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=45), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=5, max_row=45))
ch.series[0].smooth = False
ch.series[0].graphicalProperties.line.solidFill = "0E9F8E"; ch.series[0].graphicalProperties.line.width = 28000
ch.legend = None
ws.add_chart(ch, "E12")
section(ws, 4, 5, "What-if calculator", 2)
label_cell(ws, "E5", "Attainment", bold=True, fill=FILL_GREY); input_cell(ws, "F5", 1.12, fmt=FMT_PCT0)
label_cell(ws, "E6", "Target incentive", bold=True, fill=FILL_GREY); input_cell(ws, "F6", 25000, fmt=FMT_MONEY)
label_cell(ws, "E7", "Payout % of TI", bold=True, fill=FILL_GREY); formula_cell(ws, "F7", f"={payout_formula('F5')}", fmt=FMT_PCT)
label_cell(ws, "E8", "Payout amount", bold=True, fill=FILL_GREY); formula_cell(ws, "F8", "=ROUND(F7*F6,0)", fmt=FMT_MONEY)
label_cell(ws, "E9", "Kicker", bold=True, fill=FILL_GREY); formula_cell(ws, "F9", f"=IF(AND({KICK}>0,F5>={KICK_AT}),{KICK},0)", fmt=FMT_MONEY)
label_cell(ws, "E10", "Total", bold=True, fill=FILL_TEAL_LIGHT); formula_cell(ws, "F10", "=F8+F9", fmt=FMT_MONEY, bold=True, fill=FILL_TEAL_LIGHT)

# ---------------------------------------------------------------- Dashboard
ws = wb.create_sheet("Dashboard"); hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 34, "C": 16, "D": 16, "E": 16, "F": 3, "G": 30, "H": 14, "I": 14, "J": 14, "K": 3})
title_band(ws, "Dashboard", "Scenario comparison for the period in Plan Design. No inputs on this tab.", 11)
S = "Simulation"
tr = 5 + NREPS
def rng(col): return f"{S}!${col}$5:${col}${4+NREPS}"
kpi_tile(ws, "B", 4, "PLAN COST — BASE", f"={S}!$I${tr}", FMT_MONEY, 1)
kpi_tile(ws, "C", 4, "COST ÷ SALES — BASE", f"=IF({S}!$D${tr}>0,{S}!$I${tr}/{S}!$D${tr},0)", FMT_PCT, 1)
kpi_tile(ws, "D", 4, "REPS PAID — BASE", f'=COUNTIF({rng("G")},">0")', "0", 1)
kpi_tile(ws, "E", 4, "REPS AT CAP — BASE", f'=COUNTIF({rng("F")},">="&{CAP})', "0", 1)
section(ws, 7, 2, "Scenario comparison", 4)
header_row(ws, 8, 2, ["Metric", "Base", "Upside", "Downside"])
metrics = [
    ("Total target", lambda a, t, p, pay, k, tot: f"={S}!$B${tr}", FMT_MONEY),
    ("Total actual sales", lambda a, t, p, pay, k, tot: f"={S}!${a}${tr}", FMT_MONEY),
    ("Team attainment", lambda a, t, p, pay, k, tot: f"=IF({S}!$B${tr}>0,{S}!${a}${tr}/{S}!$B${tr},0)", FMT_PCT),
    ("Total target incentive", lambda a, t, p, pay, k, tot: f"={S}!$C${tr}", FMT_MONEY),
    ("Incentive payout", lambda a, t, p, pay, k, tot: f"={S}!${pay}${tr}", FMT_MONEY),
    ("Kickers", lambda a, t, p, pay, k, tot: f"={S}!${k}${tr}", FMT_MONEY),
    ("Total plan cost", lambda a, t, p, pay, k, tot: f"={S}!${tot}${tr}", FMT_MONEY),
    ("Cost ÷ target incentive", lambda a, t, p, pay, k, tot: f"=IF({S}!$C${tr}>0,{S}!${tot}${tr}/{S}!$C${tr},0)", FMT_PCT),
    ("Cost ÷ actual sales", lambda a, t, p, pay, k, tot: f"=IF({S}!${a}${tr}>0,{S}!${tot}${tr}/{S}!${a}${tr},0)", FMT_PCT),
    ("Reps paid (payout > 0)", lambda a, t, p, pay, k, tot: f'=COUNTIF({rng(pay)},">0")', "0"),
    ("Reps at or above 100%", lambda a, t, p, pay, k, tot: f'=COUNTIF({rng(t)},">=1")', "0"),
    ("Reps at cap", lambda a, t, p, pay, k, tot: f'=COUNTIF({rng(p)},">="&{CAP})', "0"),
    ("Highest individual payout", lambda a, t, p, pay, k, tot: f"=MAX({rng(tot)})", FMT_MONEY),
    ("Average payout per paid rep", lambda a, t, p, pay, k, tot: f'=IFERROR(SUMIF({rng(tot)},">0")/COUNTIF({rng(tot)},">0"),0)', FMT_MONEY),
]
for mi, (name, fn, fmt) in enumerate(metrics):
    r = 9 + mi
    label_cell(ws, f"B{r}", name, bold=True, fill=FILL_GREY)
    for si in range(3):
        c0 = 4 + si * 6
        a, t, p, pay, k, tot = [L(c0 + j) for j in range(6)]
        formula_cell(ws, f"{L(3+si)}{r}", fn(a, t, p, pay, k, tot), fmt=fmt, bold=(name == "Total plan cost"))
section(ws, 7, 7, "Attainment distribution (reps)", 4)
header_row(ws, 8, 7, ["Band", "Base", "Upside", "Downside"])
bands = [("Below threshold", None, f"INDEX({A_RNG},1)"), ("Threshold to <100%", f"INDEX({A_RNG},1)", "1"),
         ("100% to <120%", "1", "1.2"), ("120% and above", "1.2", None)]
for bi, (name, lo, hi) in enumerate(bands):
    r = 9 + bi
    label_cell(ws, f"G{r}", name, bold=True, fill=FILL_GREY)
    for si in range(3):
        t = L(4 + si * 6 + 1)
        a_col = L(4 + si * 6)
        if lo is None:
            f = f'=COUNTIF({rng(t)},"<"&{hi})'
        elif hi is None:
            f = f'=COUNTIF({rng(t)},">="&{lo})'
        else:
            f = f'=COUNTIF({rng(t)},">="&{lo})-COUNTIF({rng(t)},">="&{hi})'
        formula_cell(ws, f"{L(8+si)}{r}", f, fmt="0")
label_cell(ws, "G13", "Total reps", bold=True, fill=FILL_TEAL_LIGHT)
for si in range(3):
    formula_cell(ws, f"{L(8+si)}13", f"=SUM({L(8+si)}9:{L(8+si)}12)", fmt="0", bold=True, fill=FILL_TEAL_LIGHT)
bc = BarChart(); bc.type = "col"; bc.grouping = "clustered"; bc.title = "Attainment distribution by scenario"
bc.height, bc.width = 8, 14
bc.add_data(Reference(ws, min_col=8, min_row=8, max_col=10, max_row=12), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=7, min_row=9, max_row=12))
for s_, colr in zip(bc.series, ["0E9F8E", "1F2A44", "4B5D80"]):
    s_.graphicalProperties.solidFill = colr
ws.add_chart(bc, "G15")
note(ws, "B25", "Reading the dashboard: cost ÷ actual sales is the number finance will ask about. If Upside pushes it past your budget, "
                "lower the payout at the top points or tighten the cap on Plan Design and watch this tab update.", merge_to="E27")

for s in wb.worksheets:
    s.sheet_properties.tabColor = TEAL if s.title in ("Plan Design", "Team") else NAVY
wb["Start Here"].sheet_properties.tabColor = "F2B134"
wb.save(OUT); print("saved", OUT)
