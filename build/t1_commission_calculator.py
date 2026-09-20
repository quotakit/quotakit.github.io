"""QuotaKit — Tiered Sales Commission Calculator (Excel + Google Sheets)."""
import sys, random, datetime as dt
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter as L
from qk_style import *

SAMPLE = "--sample" in sys.argv
OUT = sys.argv[-1] if sys.argv[-1].endswith(".xlsx") else "out.xlsx"
NREPS = 50
NLOG = 1000
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

wb = Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------------- Start Here
instructions_sheet(
    wb, "Tiered Sales Commission Calculator",
    steps=[
        "Settings tab: type your currency, plan year, commission tiers (up to 5) and pick the tier method.",
        "Reps tab: list your sales reps (up to 50) with their team and monthly quota.",
        "Sales Log tab: enter each sale — date, rep (dropdown), customer, amount. Up to 1,000 rows.",
        "Monthly Sales / Monthly Commission tabs fill in automatically. Use Statement to print one rep's month.",
        "Dashboard shows year-to-date totals, the monthly trend and your top 10 reps.",
    ],
    tabs=[
        ("Settings", "Currency, plan year, tier thresholds and rates, tier method (marginal or whole-amount), quota bonus, monthly cap."),
        ("Reps", "Your rep roster: name, team, monthly quota. Names feed every dropdown."),
        ("Sales Log", "One row per sale. This is the only place you enter sales."),
        ("Monthly Sales", "Rep × month sales matrix, summed from the Sales Log."),
        ("Monthly Commission", "Rep × month commission and quota bonus, plus annual totals."),
        ("Statement", "Pick a rep and a month: shows the tier-by-tier breakdown, bonus, cap and total payout. Print-ready."),
        ("Dashboard", "KPI tiles, monthly chart, top 10 reps."),
    ],
    tips=[
        "Marginal method = each slice of sales is paid at its own tier rate (like income-tax brackets). Whole-amount method = the entire month's sales are paid at the rate of the highest tier reached.",
        "Unused tiers: leave 'From' and 'Rate' blank. Tier 1 should always start at 0.",
        "Set the quota bonus to 0 if you do not pay a flat bonus for hitting quota. Set the monthly cap to 0 for no cap.",
        "To reset: clear the yellow cells in Sales Log and Reps. Never delete formula columns.",
    ],
)

# ---------------------------------------------------------------- Settings
ws = wb.create_sheet("Settings")
hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 34, "C": 18, "D": 16, "E": 18, "F": 3, "G": 60})
title_band(ws, "Settings", "Edit the yellow cells. Everything else in the workbook reads from here.", 7)
section(ws, 4, 2, "General")
label_cell(ws, "B5", "Currency symbol / code", bold=True, fill=FILL_GREY)
input_cell(ws, "C5", "$")
label_cell(ws, "B6", "Plan year", bold=True, fill=FILL_GREY)
input_cell(ws, "C6", 2026, fmt="0")
label_cell(ws, "B7", "Tier method", bold=True, fill=FILL_GREY)
input_cell(ws, "C7", "Marginal")
add_list_validation(ws, "C7", '"Marginal,Whole amount"', "Marginal: each slice at its own rate. Whole amount: all sales at the top tier reached.")
note(ws, "G5", "Currency is a label only; it does not convert anything.")
note(ws, "G6", "Sales dated outside the plan year are ignored in the monthly matrices.")
note(ws, "G7", "Marginal works like tax brackets. Whole amount pays the whole month at the highest tier's rate.")

section(ws, 9, 2, "Commission tiers (monthly sales per rep)")
header_row(ws, 10, 2, ["Tier", "From (sales ≥)", "Rate", "Incremental rate"], widths=None)
tiers = [(0, 0.03), (25000, 0.05), (50000, 0.07), (100000, 0.10), (None, None)]
for i, (frm, rate) in enumerate(tiers, 1):
    r = 10 + i
    label_cell(ws, f"B{r}", f"Tier {i}", bold=True, fill=FILL_GREY)
    input_cell(ws, f"C{r}", frm if frm is not None else None, fmt=FMT_MONEY)
    input_cell(ws, f"D{r}", rate if rate is not None else None, fmt=FMT_PCT)
    if i == 1:
        formula_cell(ws, f"E{r}", f'=IF(OR(C{r}="",D{r}=""),0,D{r})', fmt=FMT_PCT)
    else:
        formula_cell(ws, f"E{r}", f'=IF(OR(C{r}="",D{r}=""),0,D{r}-D{r-1})', fmt=FMT_PCT)
note(ws, "G10", "'From' must increase down the table and Tier 1 must start at 0. Leave unused tiers blank. "
                "'Incremental rate' is a helper used by the marginal method — do not edit.", merge_to="G13")
ws.row_dimensions[10].height = 28
label_cell(ws, "B16", "Active tiers", bold=True, fill=FILL_GREY)
formula_cell(ws, "C16", "=COUNT(C11:C15)", fmt="0")

section(ws, 18, 2, "Quota bonus and cap")
label_cell(ws, "B19", "Quota bonus (flat, per rep per month)", bold=True, fill=FILL_GREY)
input_cell(ws, "C19", 500, fmt=FMT_MONEY)
label_cell(ws, "B20", "Bonus paid at attainment ≥", bold=True, fill=FILL_GREY)
input_cell(ws, "C20", 1.0, fmt=FMT_PCT0)
label_cell(ws, "B21", "Monthly commission cap per rep (0 = no cap)", bold=True, fill=FILL_GREY)
input_cell(ws, "C21", 0, fmt=FMT_MONEY)
note(ws, "G19", "Bonus is paid when the rep's monthly sales ÷ monthly quota reaches this attainment. Set bonus to 0 to switch it off.", merge_to="G21")

# Named-ish anchors (we use direct refs; documented here)
# Tier froms: Settings!$C$11:$C$15 ; rates: $D$11:$D$15 ; incr: $E$11:$E$15 ; method $C$7 ; year $C$6
# bonus $C$19 ; threshold $C$20 ; cap $C$21 ; active tiers $C$16

# ---------------------------------------------------------------- Reps
ws = wb.create_sheet("Reps")
hide_gridlines(ws)
set_widths(ws, {"A": 26, "B": 18, "C": 18, "D": 3, "E": 60})
title_band(ws, "Reps", "Up to 50 reps. Names here feed the dropdown in Sales Log.", 5)
header_row(ws, 4, 1, ["Rep name", "Team", "Monthly quota"])
sample_reps = [
    ("Aisha Khan", "North", 40000), ("Ben Carter", "North", 35000), ("Chloe Martin", "South", 45000),
    ("Diego Alvarez", "South", 30000), ("Emma Wilson", "West", 50000), ("Farhan Ali", "West", 30000),
    ("Grace Lee", "East", 40000), ("Hiro Tanaka", "East", 35000),
]
for i in range(NREPS):
    r = 5 + i
    if SAMPLE and i < len(sample_reps):
        n, t, q = sample_reps[i]
        input_cell(ws, f"A{r}", n)
        input_cell(ws, f"B{r}", t)
        input_cell(ws, f"C{r}", q, fmt=FMT_MONEY)
    else:
        input_cell(ws, f"A{r}")
        input_cell(ws, f"B{r}")
        input_cell(ws, f"C{r}", fmt=FMT_MONEY)
note(ws, "E4", "Tip: keep names unique. If two reps share a name, add an initial. Quota is used for attainment and the quota bonus.", merge_to="E6")
freeze(ws, "A5")

# ---------------------------------------------------------------- Sales Log
ws = wb.create_sheet("Sales Log")
hide_gridlines(ws)
set_widths(ws, {"A": 14, "B": 24, "C": 26, "D": 16, "E": 30, "F": 12})
title_band(ws, "Sales Log", "One row per sale. Date, rep (dropdown), customer, amount. Up to 1,000 rows.", 6)
header_row(ws, 4, 1, ["Date", "Rep", "Customer", "Amount", "Notes", "Month key"])
add_list_validation(ws, f"B5:B{4+NLOG}", f"=Reps!$A$5:$A${4+NREPS}", "Pick the rep from the Reps tab.")
random.seed(7)
customers = ["Northwind Ltd", "Acme Corp", "Globex", "Initech", "Umbrella Co", "Stark Industries", "Wayne Enterprises",
             "Hooli", "Vandelay Imports", "Pied Piper", "Soylent Inc", "Wonka Foods", "Dunder Mifflin", "Cyberdyne"]
rows = []
if SAMPLE:
    for m in range(1, 10):  # Jan–Sep
        for n, t, q in sample_reps:
            k = random.randint(3, 7)
            for _ in range(k):
                d = dt.date(2026, m, random.randint(1, 28))
                amt = random.choice([1500, 2200, 3400, 4800, 6000, 7500, 9800, 12500, 15000])
                rows.append((d, n, random.choice(customers), amt))
    rows.sort()
for i in range(NLOG):
    r = 5 + i
    if i < len(rows):
        d, n, c, a = rows[i]
        input_cell(ws, f"A{r}", d, fmt=FMT_DATE)
        input_cell(ws, f"B{r}", n)
        input_cell(ws, f"C{r}", c)
        input_cell(ws, f"D{r}", a, fmt=FMT_MONEY)
        input_cell(ws, f"E{r}")
    else:
        input_cell(ws, f"A{r}", fmt=FMT_DATE)
        input_cell(ws, f"B{r}")
        input_cell(ws, f"C{r}")
        input_cell(ws, f"D{r}", fmt=FMT_MONEY)
        input_cell(ws, f"E{r}")
    formula_cell(ws, f"F{r}", f'=IF(A{r}="","",YEAR(A{r})&"-"&TEXT(MONTH(A{r}),"00"))', align=CENTER)
    ws[f"F{r}"].font = F_NOTE
freeze(ws, "A5")

# ---------------------------------------------------------------- Monthly Sales
ws = wb.create_sheet("Monthly Sales")
hide_gridlines(ws)
set_widths(ws, {"A": 26})
for c in range(2, 15):
    ws.column_dimensions[L(c)].width = 12
title_band(ws, "Monthly Sales", "Rep × month sales, summed from the Sales Log for the plan year. No inputs on this tab.", 14)
header_row(ws, 4, 1, ["Rep"] + MONTHS + ["Total"])
# month key row (hidden helper) in row 3
for m in range(12):
    col = L(2 + m)
    ws[f"{col}3"] = f'=Settings!$C$6&"-"&TEXT({m+1},"00")'
    ws[f"{col}3"].font = F_NOTE
    ws[f"{col}3"].alignment = CENTER
ws["A3"] = "month key →"
ws["A3"].font = F_NOTE
for i in range(NREPS):
    r = 5 + i
    formula_cell(ws, f"A{r}", f'=IF(Reps!A{r}="","",Reps!A{r})', align=LEFT)
    for m in range(12):
        col = L(2 + m)
        formula_cell(ws, f"{col}{r}",
                     f'=IF($A{r}="","",SUMIFS(\'Sales Log\'!$D$5:$D${4+NLOG},\'Sales Log\'!$B$5:$B${4+NLOG},$A{r},\'Sales Log\'!$F$5:$F${4+NLOG},{col}$3))',
                     fmt=FMT_MONEY)
    formula_cell(ws, f"N{r}", f'=IF($A{r}="","",SUM(B{r}:M{r}))', fmt=FMT_MONEY, bold=True)
tr = 5 + NREPS
label_cell(ws, f"A{tr}", "Total", bold=True, fill=FILL_GREY)
for m in range(13):
    col = L(2 + m)
    formula_cell(ws, f"{col}{tr}", f"=SUM({col}5:{col}{tr-1})", fmt=FMT_MONEY, bold=True, fill=FILL_GREY)
freeze(ws, "B5")

# ---------------------------------------------------------------- Monthly Commission
ws = wb.create_sheet("Monthly Commission")
hide_gridlines(ws)
set_widths(ws, {"A": 26})
for c in range(2, 32):
    ws.column_dimensions[L(c)].width = 12
title_band(ws, "Monthly Commission", "Commission (B–M), quota bonus (O–Z) and annual totals (AB–AD). No inputs on this tab.", 30)
ws["B3"] = "Commission by month"; ws["B3"].font = F_SECTION
ws["O3"] = "Quota bonus by month"; ws["O3"].font = F_SECTION
ws["AB3"] = "Annual totals"; ws["AB3"].font = F_SECTION
header_row(ws, 4, 1, ["Rep"] + MONTHS)
header_row(ws, 4, 15, MONTHS, fill=FILL_NAVY)
header_row(ws, 4, 28, ["Commission", "Bonus", "Total payout"], fill=FILL_TEAL)
ws.column_dimensions["N"].width = 3
ws.column_dimensions["AA"].width = 3
FR, RT, IN = "Settings!$C$11:$C$15", "Settings!$D$11:$D$15", "Settings!$E$11:$E$15"
for i in range(NREPS):
    r = 5 + i
    formula_cell(ws, f"A{r}", f'=IF(Reps!A{r}="","",Reps!A{r})', align=LEFT)
    for m in range(12):
        col = L(2 + m)
        S = f"'Monthly Sales'!{col}{r}"
        marginal = f"SUMPRODUCT(({S}>{FR})*({S}-{FR})*{IN})"
        whole = f"IFERROR(INDEX({RT},MATCH({S},INDEX({FR},1):INDEX({FR},Settings!$C$16),1))*{S},0)"
        raw = f'IF(Settings!$C$7="Marginal",{marginal},{whole})'
        capped = f"IF(Settings!$C$21>0,MIN({raw},Settings!$C$21),{raw})"
        formula_cell(ws, f"{col}{r}", f'=IF($A{r}="","",{capped})', fmt=FMT_MONEY)
        bcol = L(15 + m)
        Q = f"Reps!$C{r}"
        formula_cell(ws, f"{bcol}{r}",
                     f'=IF($A{r}="","",IF(AND(Settings!$C$19>0,{Q}>0),IF({S}/{Q}>=Settings!$C$20,Settings!$C$19,0),0))',
                     fmt=FMT_MONEY)
    formula_cell(ws, f"AB{r}", f'=IF($A{r}="","",SUM(B{r}:M{r}))', fmt=FMT_MONEY, bold=True)
    formula_cell(ws, f"AC{r}", f'=IF($A{r}="","",SUM(O{r}:Z{r}))', fmt=FMT_MONEY, bold=True)
    formula_cell(ws, f"AD{r}", f'=IF($A{r}="","",AB{r}+AC{r})', fmt=FMT_MONEY, bold=True, fill=FILL_TEAL_LIGHT)
    # rank helper (hidden-ish) in AF: YTD sales with tie-break
    formula_cell(ws, f"AF{r}", f'=IF($A{r}="",-1,\'Monthly Sales\'!N{r}+ROW()/1000000)', fmt="0.000000")
    ws[f"AF{r}"].font = F_NOTE
ws["AF4"] = "rank key"; ws["AF4"].font = F_NOTE
tr = 5 + NREPS
label_cell(ws, f"A{tr}", "Total", bold=True, fill=FILL_GREY)
for c in list(range(2, 14)) + list(range(15, 27)) + [28, 29, 30]:
    col = L(c)
    formula_cell(ws, f"{col}{tr}", f"=SUM({col}5:{col}{tr-1})", fmt=FMT_MONEY, bold=True, fill=FILL_GREY)
freeze(ws, "B5")

# ---------------------------------------------------------------- Statement
ws = wb.create_sheet("Statement")
hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 30, "C": 18, "D": 18, "E": 18, "F": 18, "G": 3})
title_band(ws, "Commission Statement", "Pick a rep and a month. Print this tab for the rep.", 7)
label_cell(ws, "B4", "Rep", bold=True, fill=FILL_GREY)
input_cell(ws, "C4", sample_reps[0][0] if SAMPLE else None)
add_list_validation(ws, "C4", f"=Reps!$A$5:$A${4+NREPS}")
label_cell(ws, "B5", "Month (1–12)", bold=True, fill=FILL_GREY)
input_cell(ws, "C5", 6 if SAMPLE else 1, fmt="0")
add_list_validation(ws, "C5", '"1,2,3,4,5,6,7,8,9,10,11,12"')
label_cell(ws, "B6", "Plan year", bold=True, fill=FILL_GREY)
formula_cell(ws, "C6", "=Settings!$C$6", fmt="0")
label_cell(ws, "B7", "Team", bold=True, fill=FILL_GREY)
formula_cell(ws, "C7", f'=IFERROR(INDEX(Reps!$B$5:$B${4+NREPS},MATCH($C$4,Reps!$A$5:$A${4+NREPS},0)),"")', align=LEFT)

section(ws, 9, 2, "Summary")
label_cell(ws, "B10", "Sales this month", bold=True, fill=FILL_GREY)
formula_cell(ws, "C10", f"=IF($C$4=\"\",0,IFERROR(N(INDEX('Monthly Sales'!$B$5:$M${4+NREPS},MATCH($C$4,'Monthly Sales'!$A$5:$A${4+NREPS},0),$C$5)),0))", fmt=FMT_MONEY, bold=True)
label_cell(ws, "B11", "Monthly quota", bold=True, fill=FILL_GREY)
formula_cell(ws, "C11", f"=IF($C$4=\"\",0,IFERROR(N(INDEX(Reps!$C$5:$C${4+NREPS},MATCH($C$4,Reps!$A$5:$A${4+NREPS},0))),0))", fmt=FMT_MONEY)
label_cell(ws, "B12", "Quota attainment", bold=True, fill=FILL_GREY)
formula_cell(ws, "C12", "=IF(C11>0,C10/C11,0)", fmt=FMT_PCT)
label_cell(ws, "B13", "Tier method", bold=True, fill=FILL_GREY)
formula_cell(ws, "C13", "=Settings!$C$7", align=LEFT)
label_cell(ws, "B14", "Highest tier reached", bold=True, fill=FILL_GREY)
formula_cell(ws, "C14", f'=IFERROR("Tier "&MATCH(C10,INDEX({FR},1):INDEX({FR},Settings!$C$16),1),"—")', align=LEFT)

section(ws, 16, 2, "Tier-by-tier breakdown")
header_row(ws, 17, 2, ["Tier", "From", "To", "Sales in tier", "Rate", "Commission"])
ws.column_dimensions["G"].width = 18
for i in range(5):
    r = 18 + i
    sr = 11 + i
    label_cell(ws, f"B{r}", f"Tier {i+1}", bold=True, fill=FILL_GREY)
    formula_cell(ws, f"C{r}", f'=IF(Settings!C{sr}="","",Settings!C{sr})', fmt=FMT_MONEY)
    nxt = f"Settings!C{sr+1}" if i < 4 else '""'
    formula_cell(ws, f"D{r}", f'=IF(Settings!C{sr}="","",IF({nxt}="","and above",{nxt}))' if i < 4 else f'=IF(Settings!C{sr}="","","and above")', fmt=FMT_MONEY)
    # sales slice in this tier
    upper = f"IF({nxt}=\"\",$C$10,MIN($C$10,{nxt}))" if i < 4 else "$C$10"
    formula_cell(ws, f"E{r}", f'=IF(Settings!C{sr}="","",MAX(0,{upper}-Settings!C{sr}))', fmt=FMT_MONEY)
    formula_cell(ws, f"F{r}", f'=IF(Settings!D{sr}="","",Settings!D{sr})', fmt=FMT_PCT)
    formula_cell(ws, f"G{r}", f'=IF(Settings!C{sr}="","",IF($C$13="Marginal",E{r}*F{r},IF($C$14="Tier {i+1}",$C$10*F{r},0)))', fmt=FMT_MONEY)
label_cell(ws, "B23", "Commission before cap", bold=True, fill=FILL_GREY)
formula_cell(ws, "G23", "=SUM(G18:G22)", fmt=FMT_MONEY, bold=True)
label_cell(ws, "B24", "Monthly cap applied", bold=True, fill=FILL_GREY)
formula_cell(ws, "G24", "=IF(Settings!$C$21>0,MIN(G23,Settings!$C$21),G23)", fmt=FMT_MONEY)
label_cell(ws, "B25", "Quota bonus", bold=True, fill=FILL_GREY)
formula_cell(ws, "G25", "=IF(AND(Settings!$C$19>0,C11>0),IF(C12>=Settings!$C$20,Settings!$C$19,0),0)", fmt=FMT_MONEY)
label_cell(ws, "B26", "TOTAL PAYOUT", bold=True, fill=FILL_TEAL_LIGHT)
formula_cell(ws, "G26", "=G24+G25", fmt=FMT_MONEY, bold=True, fill=FILL_TEAL_LIGHT)
for r in range(23, 27):
    ws.merge_cells(f"B{r}:F{r}")
note(ws, "B28", "Whole-amount method: only the highest tier reached shows a commission (the whole month's sales × that tier's rate). "
                "Marginal method: each tier shows the slice of sales that falls in it × its rate.", merge_to="G29")
ws.row_dimensions[28].height = 30
ws.print_area = "A1:G29"
ws.page_setup.fitToWidth = 1

# ---------------------------------------------------------------- Dashboard
ws = wb.create_sheet("Dashboard")
hide_gridlines(ws)
set_widths(ws, {"A": 3, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14, "I": 14, "J": 3, "K": 26, "L": 14, "M": 14, "N": 14})
title_band(ws, "Dashboard", "Year-to-date view of the plan year set in Settings. No inputs on this tab.", 14)
MS_TOT = f"'Monthly Sales'!$N${5+NREPS}"
MC_TOT = f"'Monthly Commission'!$AB${5+NREPS}"
MB_TOT = f"'Monthly Commission'!$AC${5+NREPS}"
kpi_tile(ws, "B", 4, "YTD SALES", f"={MS_TOT}", FMT_MONEY)
kpi_tile(ws, "D", 4, "YTD COMMISSION", f"={MC_TOT}", FMT_MONEY)
kpi_tile(ws, "F", 4, "YTD QUOTA BONUS", f"={MB_TOT}", FMT_MONEY)
kpi_tile(ws, "H", 4, "EFFECTIVE RATE", f"=IF({MS_TOT}>0,({MC_TOT}+{MB_TOT})/{MS_TOT},0)", FMT_PCT)
# monthly table for chart
section(ws, 7, 2, "Monthly trend", 6)
header_row(ws, 8, 2, ["Month", "Sales", "Commission", "Bonus", "Payout", "Eff. rate"])
for m in range(12):
    r = 9 + m
    col = L(2 + m)
    label_cell(ws, f"B{r}", MONTHS[m], fill=FILL_GREY, bold=True)
    formula_cell(ws, f"C{r}", f"='Monthly Sales'!{col}{5+NREPS}", fmt=FMT_MONEY)
    formula_cell(ws, f"D{r}", f"='Monthly Commission'!{col}{5+NREPS}", fmt=FMT_MONEY)
    formula_cell(ws, f"E{r}", f"='Monthly Commission'!{L(15+m)}{5+NREPS}", fmt=FMT_MONEY)
    formula_cell(ws, f"F{r}", f"=D{r}+E{r}", fmt=FMT_MONEY, bold=True)
    formula_cell(ws, f"G{r}", f"=IF(C{r}>0,F{r}/C{r},0)", fmt=FMT_PCT)
chart = BarChart()
chart.type = "col"
chart.title = "Sales vs payout by month"
chart.y_axis.title = "Amount"
chart.height, chart.width = 8, 18
data = Reference(ws, min_col=3, min_row=8, max_col=3, max_row=20)
cats = Reference(ws, min_col=2, min_row=9, max_row=20)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
line = LineChart()
ldata = Reference(ws, min_col=6, min_row=8, max_col=6, max_row=20)
line.add_data(ldata, titles_from_data=True)
line.y_axis.axId = 200
line.y_axis.title = "Payout"
line.y_axis.crosses = "max"
line.series[0].smooth = False
line.series[0].graphicalProperties.line.solidFill = "1F2A44"
line.series[0].graphicalProperties.line.width = 28000
chart.series[0].graphicalProperties.solidFill = "0E9F8E"
chart.series[0].graphicalProperties.line.solidFill = "0E9F8E"
chart += line
ws.add_chart(chart, "B22")
# top 10
section(ws, 7, 11, "Top 10 reps (YTD sales)", 4)
header_row(ws, 8, 11, ["Rep", "YTD sales", "Commission", "Payout"])
KEY = f"'Monthly Commission'!$AF$5:$AF${4+NREPS}"
for k in range(10):
    r = 9 + k
    formula_cell(ws, f"K{r}", f'=IFERROR(IF(LARGE({KEY},{k+1})<0,"",INDEX(\'Monthly Commission\'!$A$5:$A${4+NREPS},MATCH(LARGE({KEY},{k+1}),{KEY},0))),"")', align=LEFT)
    formula_cell(ws, f"L{r}", f'=IF(K{r}="","",INDEX(\'Monthly Sales\'!$N$5:$N${4+NREPS},MATCH(K{r},\'Monthly Sales\'!$A$5:$A${4+NREPS},0)))', fmt=FMT_MONEY)
    formula_cell(ws, f"M{r}", f'=IF(K{r}="","",INDEX(\'Monthly Commission\'!$AB$5:$AB${4+NREPS},MATCH(K{r},\'Monthly Commission\'!$A$5:$A${4+NREPS},0)))', fmt=FMT_MONEY)
    formula_cell(ws, f"N{r}", f'=IF(K{r}="","",INDEX(\'Monthly Commission\'!$AD$5:$AD${4+NREPS},MATCH(K{r},\'Monthly Commission\'!$A$5:$A${4+NREPS},0)))', fmt=FMT_MONEY, bold=True)
ws.conditional_formatting.add("L9:L18", DataBarRule(start_type="min", end_type="max", color="0E9F8E"))

for s in wb.worksheets:
    s.sheet_properties.tabColor = TEAL if s.title in ("Settings", "Reps", "Sales Log") else NAVY
wb["Start Here"].sheet_properties.tabColor = "F2B134"
wb.save(OUT)
print("saved", OUT)
