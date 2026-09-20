"""QuotaKit build pipeline (runs in GitHub Actions).

Builds the five spreadsheet products (blank + sample), recalculates them with LibreOffice so cached
values exist, renders sheets to PNG, composes listing images, packages ZIPs with read-me files, and
builds the static site. Everything lands under out/.
"""
import os, sys, subprocess, shutil, zipfile, tempfile, glob
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = Path(os.environ.get("QK_OUT", HERE.parent / "out")).resolve()
PRODUCTS = OUT / "products"; IMAGES = OUT / "images"; PACK = OUT / "packages"; SITE = OUT / "site"
for d in (PRODUCTS, IMAGES / "render", IMAGES / "listing", PACK, SITE):
    d.mkdir(parents=True, exist_ok=True)

TEMPLATES = [
    ("t1_commission_calculator", "Sales-Commission-Calculator", "Sales Commission Calculator", ["Dashboard", "Statement", "Settings"]),
    ("t2_incentive_simulator", "Incentive-Plan-Simulator", "Sales Incentive Plan Simulator", ["Dashboard", "Payout Curve", "Plan Design"]),
    ("t3_kpi_tracker", "Sales-KPI-Tracker", "Sales KPI Dashboard & Monthly Tracker", ["Dashboard", "KPI Calc", "Monthly Input"]),
    ("t4_pipeline_tracker", "Pipeline-Deal-Tracker", "Sales Pipeline & Deal Tracker", ["Dashboard", "Deals", "Settings"]),
    ("t5_coverage_planner", "Territory-Coverage-Planner", "Territory & Account Coverage Planner", ["Dashboard", "Rep Summary", "Accounts"]),
]

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def sh(cmd, **kw):
    print("+", " ".join(str(c) for c in cmd), flush=True)
    return subprocess.run([str(c) for c in cmd], check=True, **kw)


def soffice_profile():
    prof = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    url = prof.as_uri()
    subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"], capture_output=True, timeout=120)
    mdir = prof / "user" / "basic" / "Standard"
    mdir.mkdir(parents=True, exist_ok=True)
    (mdir / "Module1.xba").write_text(MACRO)
    return url


def recalc(path, url):
    subprocess.run(["timeout", "180", "soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
                    "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application", str(path)],
                   check=True, capture_output=True)


def render(prod, sheets):
    from openpyxl import load_workbook
    src = PRODUCTS / f"QuotaKit_{prod}_SAMPLE.xlsx"
    for sh_ in sheets:
        wb = load_workbook(src, data_only=True)
        for s in list(wb.sheetnames):
            if s != sh_:
                del wb[s]
        ws = wb[sh_]
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if sh_ in ("Deals", "Accounts", "KPI Calc", "Monthly Input"):
            ws.print_area = f"A1:{'V' if sh_ == 'Accounts' else 'O'}40"
        tmp = IMAGES / "render" / f"{prod}__{sh_.replace(' ', '_')}.xlsx"
        wb.save(tmp)
        subprocess.run(["timeout", "180", "soffice", "--headless", "--convert-to", "pdf", "--outdir", str(IMAGES / "render"), str(tmp)], capture_output=True)
        pdf = tmp.with_suffix(".pdf")
        subprocess.run(["pdftoppm", "-png", "-r", "170", "-f", "1", "-l", "1", str(pdf), str(tmp.with_suffix(""))], check=True)


README = """QUOTAKIT — {name}
================================================

Thank you for your purchase.

FILES IN THIS ZIP
  1. QuotaKit_{key}_BLANK.xlsx   — the template, ready for your data
  2. QuotaKit_{key}_SAMPLE.xlsx  — the same template filled with example data
  3. READ-ME-FIRST.txt            — this file

GETTING STARTED (2 minutes)
  Excel:          open the BLANK file. Start on the "Start Here" tab.
  Google Sheets:  Google Drive > New > File upload > pick the .xlsx.
                  Then open it and choose File > Save as Google Sheets.
  Numbers / LibreOffice: open directly.

  Yellow cells with blue text are yours to edit. Everything else calculates.
  No macros, so there is nothing to enable.

IF SOMETHING LOOKS WRONG
  - A column of "-" or 0 usually means an input tab is still empty.
  - Dropdowns pull from the Settings / Reps tabs — fill those first.
  - Still stuck? Reply to your purchase email and include the tab name.
    If a formula does not behave as described in the listing, we fix it or refund you.

LICENCE
  Personal and small-business use: use it inside your own company or for
  your own clients. Please do not resell, share or redistribute the file.

Created with AI assistance; designed, curated and tested by QuotaKit.
"""


def package():
    for _, key, name, _ in TEMPLATES:
        with zipfile.ZipFile(PACK / f"QuotaKit_{key}.zip", "w", zipfile.ZIP_DEFLATED) as z:
            for v in ("BLANK", "SAMPLE"):
                z.write(PRODUCTS / f"QuotaKit_{key}_{v}.xlsx", f"QuotaKit_{key}_{v}.xlsx")
            z.writestr("READ-ME-FIRST.txt", README.format(name=name, key=key))
    with zipfile.ZipFile(PACK / "QuotaKit_Sales-Ops-Toolkit_BUNDLE.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for _, key, name, _ in TEMPLATES:
            for v in ("BLANK", "SAMPLE"):
                z.write(PRODUCTS / f"QuotaKit_{key}_{v}.xlsx", f"{key}/QuotaKit_{key}_{v}.xlsx")
            z.writestr(f"{key}/READ-ME-FIRST.txt", README.format(name=name, key=key))


def covers():
    """Small JPEG covers (Gumroad) from the main listing images."""
    from PIL import Image
    cov = IMAGES / "covers"; cov.mkdir(exist_ok=True)
    for p in (IMAGES / "listing").glob("*.png"):
        im = Image.open(p).convert("RGB"); im.thumbnail((1200, 1200))
        im.save(cov / (p.stem + ".jpg"), quality=82, optimize=True)


def main():
    url = soffice_profile()
    for script, key, name, sheets in TEMPLATES:
        for flag, ver in (("--sample", "SAMPLE"), ("", "BLANK")):
            out = PRODUCTS / f"QuotaKit_{key}_{ver}.xlsx"
            args = [sys.executable, str(HERE / f"{script}.py")] + ([flag] if flag else []) + [str(out)]
            sh(args, cwd=HERE)
            recalc(out, url)
        render(key, sheets)
    env = dict(os.environ, QK_IMAGES=str(IMAGES))
    sh([sys.executable, str(HERE / "make_mockups.py")], env=env)
    package(); covers()
    env = dict(os.environ, QK_DIST=str(SITE), QK_IMG_SRC=str(IMAGES / "listing"), QK_NO_ARTIFACT="1")
    sh([sys.executable, str(HERE / "build_site.py")], env=env)
    (SITE / ".nojekyll").write_text("\n")
    print("\nOUT:")
    for p in sorted(OUT.rglob("*")):
        if p.is_file():
            print(f"  {p.relative_to(OUT)}  {p.stat().st_size:,}")


if __name__ == "__main__":
    main()
